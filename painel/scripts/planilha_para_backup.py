"""Converte a planilha de despesas da MB2 (.xlsx) num arquivo de backup do Painel MB².

Uso:  python3 scripts/planilha_para_backup.py "PLANILHA DESPESAS OBRAS MB2.xlsx" saida.json
Depois, no painel: Ajustes › Restaurar um backup › escolher o .json gerado.

Cada aba de obra tem uma tabela do Excel com DATA, DESCRIÇÃO, CATEGORIA, VALOR, OBSERVAÇÃO.
A aba RESUMO dá o nome e o tipo (Obra, Terreno, Projeto) de cada aba.
Os ids são fixos (derivados do conteúdo), então importar de novo não duplica nada.
Precisa de: pip install openpyxl
"""
import datetime as dt
import json
import re
import sys
import uuid

import openpyxl

CATEGORIAS = {
    'Terreno e aquisição', 'Projetos e documentação', 'Impostos e taxas', 'Materiais (M)', 'Mão de obra',
    'Serviços (S)', 'Locação e consumo', 'Consórcio', 'A classificar',
}
STATUS_POR_TIPO = {'Obra': 'em_andamento', 'Terreno': 'planejamento', 'Projeto': 'planejamento'}
MINUSCULAS = {'de', 'da', 'do', 'das', 'dos', 'e'}
BASE = uuid.UUID('6f1c2a5e-3b9d-4c1e-9a77-0b2d5e8f4a10')


def titulo(texto):
    palavras = []
    for i, p in enumerate(texto.split()):
        if i > 0 and p.lower() in MINUSCULAS:
            palavras.append(p.lower())
        elif p in ('JD', '–', '-') or re.fullmatch(r'[A-Z]', p) or any(c.isdigit() for c in p):
            palavras.append(p)
        else:
            palavras.append(p[:1].upper() + p[1:].lower())
    return ' '.join(palavras)


def ler_resumo(wb):
    # Linhas do RESUMO: =HYPERLINK("#'ABA'!A1","NOME DA OBRA") | TIPO
    obras = []
    for link, tipo in wb['RESUMO'].iter_rows(min_row=4, max_col=2, values_only=True):
        m = re.match(r'=HYPERLINK\("#\'(.+?)\'!A1","(.+?)"\)', str(link or ''))
        if m:
            obras.append((m.group(1), m.group(2), tipo))
    return obras


def ler_tabela(ws):
    tabela = next(iter(ws.tables.values()))
    linhas = list(ws[tabela.ref])
    cab = [c.value for c in linhas[0]]
    idx = {nome: cab.index(nome) for nome in ('DATA', 'DESCRIÇÃO', 'CATEGORIA', 'VALOR', 'OBSERVAÇÃO')}
    for linha in linhas[1:]:
        v = [c.value for c in linha]
        if all(x is None for x in v):
            continue
        yield {k: v[i] for k, i in idx.items()}


def notas_fora_da_tabela(ws):
    # Texto solto abaixo da tabela (ex.: "OBSERVAÇÕES" + nota) vira observação da obra.
    fim = int(re.search(r'(\d+)$', next(iter(ws.tables.values())).ref).group(1))
    textos = [r[0] for r in ws.iter_rows(min_row=fim + 1, max_col=1, values_only=True) if isinstance(r[0], str)]
    return [t for t in textos if t.strip().upper() != 'OBSERVAÇÕES']


def converter(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    links = openpyxl.load_workbook(caminho)  # sem data_only, para ler as fórmulas HYPERLINK
    agora = dt.datetime.now(dt.timezone.utc).isoformat().replace('+00:00', 'Z')
    obras, custos, avisos = [], [], []

    for aba, nome, tipo in ler_resumo(links):
        ws = wb[aba]
        obra_id = str(uuid.uuid5(BASE, 'obra:' + aba))
        observacoes = [f'Tipo na planilha: {tipo}.'] + notas_fora_da_tabela(ws)
        obras.append({
            'id': obra_id, 'nome': titulo(nome), 'endereco': '', 'cliente': '', 'orcamento': 0, 'areaM2': None,
            'inicio': None, 'previsaoFim': None, 'status': STATUS_POR_TIPO.get(tipo, 'planejamento'),
            'observacoes': '\n'.join(observacoes), 'criadoEm': agora, 'atualizadoEm': agora,
        })
        vistos = {}
        for n, r in enumerate(ler_tabela(ws), start=1):
            data, valor, categoria = r['DATA'], r['VALOR'], r['CATEGORIA'] or 'A classificar'
            if not isinstance(data, (dt.date, dt.datetime)) or not isinstance(valor, (int, float)):
                avisos.append(f'{aba}, linha {n}: sem data ou valor, ignorada ({r})')
                continue
            if categoria not in CATEGORIAS:
                avisos.append(f'{aba}, linha {n}: categoria "{categoria}" desconhecida, foi para "A classificar"')
                categoria = 'A classificar'
            descricao = str(r['DESCRIÇÃO'] or '').strip()
            fornecedor = ''
            m = re.match(r'MÃO DE OBRA\s*[–-]\s*(.+)', descricao)
            if m:
                fornecedor = titulo(m.group(1))
            if r['OBSERVAÇÃO']:
                descricao = f"{descricao} ({str(r['OBSERVAÇÃO']).strip()})"
            chave = f'{aba}|{data:%Y-%m-%d}|{descricao}|{valor}'
            vistos[chave] = vistos.get(chave, 0) + 1  # lançamentos iguais no mesmo dia são legítimos
            custos.append({
                'id': str(uuid.uuid5(BASE, f'{chave}|{vistos[chave]}')), 'obraId': obra_id,
                'data': f'{data:%Y-%m-%d}', 'categoria': categoria, 'descricao': descricao, 'fornecedor': fornecedor,
                'valor': round(valor * 100), 'pagamento': 'outro', 'pago': True, 'criadoEm': agora, 'atualizadoEm': agora,
            })

    backup = {'app': 'mb2-painel', 'versao': 1, 'geradoEm': agora, 'obras': obras, 'custos': custos}
    return backup, avisos


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    backup, avisos = converter(sys.argv[1])
    with open(sys.argv[2], 'w', encoding='utf-8') as f:
        json.dump(backup, f, ensure_ascii=False, indent=1)
    for a in avisos:
        print('Aviso:', a)
    total = sum(c['valor'] for c in backup['custos']) / 100
    print(f"{len(backup['obras'])} obras, {len(backup['custos'])} lançamentos, total R$ {total:,.2f}")
    for o in backup['obras']:
        t = sum(c['valor'] for c in backup['custos'] if c['obraId'] == o['id']) / 100
        print(f"  {o['nome']}: R$ {t:,.2f}")
